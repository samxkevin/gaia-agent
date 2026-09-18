from pathlib import Path

from smolagents import Tool

from config import MAX_TOOL_TEXT


class ReadFileTool(Tool):
    name = "read_local_file"
    description = (
        "Read a local GAIA attachment. Supports UTF 8 text, JSON, CSV, XML, "
        "HTML, Markdown, Python, YAML, PDF, DOCX, and XLSX files."
    )
    inputs = {
        "path": {
            "type": "string",
            "description": "Absolute or repository local path to the file.",
        },
    }
    output_type = "string"

    def forward(self, path: str) -> str:
        file_path = Path(path).expanduser().resolve()

        if not file_path.is_file():
            return f"File not found: {file_path}"

        suffix = file_path.suffix.lower()

        try:
            if suffix == ".pdf":
                from pypdf import PdfReader

                reader = PdfReader(str(file_path))
                pages = [page.extract_text() or "" for page in reader.pages]
                return "\n\n".join(pages)[:MAX_TOOL_TEXT]

            if suffix == ".docx":
                from docx import Document

                document = Document(str(file_path))
                parts = [paragraph.text for paragraph in document.paragraphs]
                for table in document.tables:
                    for row in table.rows:
                        parts.append(" | ".join(cell.text for cell in row.cells))
                return "\n".join(parts)[:MAX_TOOL_TEXT]

            if suffix == ".xlsx":
                from openpyxl import load_workbook

                workbook = load_workbook(
                    str(file_path),
                    read_only=True,
                    data_only=True,
                )
                parts = []
                for sheet in workbook.worksheets:
                    parts.append(f"[Sheet: {sheet.title}]")
                    for row in sheet.iter_rows(values_only=True):
                        values = ["" if value is None else str(value) for value in row]
                        parts.append(" | ".join(values))
                workbook.close()
                return "\n".join(parts)[:MAX_TOOL_TEXT]

            text_suffixes = {
                ".txt", ".md", ".csv", ".json", ".xml", ".html",
                ".htm", ".py", ".js", ".ts", ".yaml", ".yml",
            }

            if suffix in text_suffixes:
                return file_path.read_text(
                    encoding="utf-8",
                    errors="replace",
                )[:MAX_TOOL_TEXT]

        except Exception as exc:
            return f"Could not read {file_path.name}: {exc}"

        return f"Unsupported readable format: {suffix}"


class InspectFileTool(Tool):
    name = "inspect_local_file"
    description = (
        "Inspect a local GAIA attachment and return path, type, size, "
        "and basic metadata."
    )
    inputs = {
        "path": {
            "type": "string",
            "description": "Absolute or repository local path to the file.",
        },
    }
    output_type = "string"

    def forward(self, path: str) -> str:
        file_path = Path(path).expanduser().resolve()

        if not file_path.is_file():
            return f"File not found: {file_path}"

        result = [
            f"path: {file_path}",
            f"name: {file_path.name}",
            f"extension: {file_path.suffix.lower()}",
            f"size_bytes: {file_path.stat().st_size}",
        ]

        if file_path.suffix.lower() in {
            ".png", ".jpg", ".jpeg", ".webp", ".gif", ".bmp",
        }:
            try:
                from PIL import Image

                with Image.open(file_path) as image:
                    result.append(f"format: {image.format}")
                    result.append(f"mode: {image.mode}")
                    result.append(f"dimensions: {image.size[0]}x{image.size[1]}")
            except Exception as exc:
                result.append(f"image_error: {exc}")

        if file_path.suffix.lower() == ".xlsx":
            try:
                from openpyxl import load_workbook

                workbook = load_workbook(
                    str(file_path),
                    read_only=True,
                    data_only=False,
                )
                result.append(f"sheets: {', '.join(workbook.sheetnames)}")
                workbook.close()
            except Exception as exc:
                result.append(f"xlsx_error: {exc}")

        return "\n".join(result)
